import csv
import os
import io
import re
import traceback
import zipfile
import urllib
import boto3
from datetime import timezone, datetime, timedelta
from typing import Dict, Tuple, List, Any
from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from sqlalchemy import select, and_
from sqlalchemy.ext.asyncio import AsyncSession
from azure.storage.blob import BlobServiceClient
from botocore.exceptions import ClientError

from app.core.utils.db_utils import *
from app.core.config import get_settings
from app.schemas.logger import logger


async def report_bulk_download(session_id: str):
    try:
        print("bulk check 1")
        storage_url = get_settings().storage.storage_account_url
        container_name = get_settings().storage.container_name
        sas_token = str(get_settings().storage.sas_token)
        print("bulk check 2")

        blob_service_client = BlobServiceClient(
            account_url=storage_url,
            credential=sas_token
        )
        container_client = blob_service_client.get_container_client(container_name)

        logger.info(f"[BULK DEBUG] container_name = {container_name}")
        logger.info(f"[BULK DEBUG] session_id = {session_id}")

        session_prefix = f"{session_id}/"
        logger.info(f"[BULK DEBUG] session_prefix = {session_prefix}")

        blob_list = list(container_client.list_blobs(name_starts_with=session_prefix))

        logger.info(f"[BULK DEBUG] blob count = {len(blob_list)}")
        for blob in blob_list:
            logger.info(f"[BULK DEBUG] blob found = {blob.name}")

        if not blob_list:
            raise HTTPException(
                status_code=404,
                detail=f"No files found for session_id {session_id}"
            )

        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            for blob in blob_list:
                blob_name = blob.name

                if blob_name.endswith("/"):
                    logger.info(f"[BULK DEBUG] skipping folder blob = {blob_name}")
                    continue

                logger.info(f"[BULK DEBUG] downloading blob = {blob_name}")

                blob_client = blob_service_client.get_blob_client(
                    container=container_name,
                    blob=blob_name
                )
                file_data = blob_client.download_blob().readall()

                logger.info(f"[BULK DEBUG] downloaded bytes = {len(file_data)} for blob = {blob_name}")
                logger.info(f"[BULK DEBUG] first 8 bytes = {file_data[:8]}")

                relative_path = blob_name[len(session_prefix):]
                logger.info(f"[BULK DEBUG] writing zip entry = {relative_path}")

                zip_file.writestr(relative_path, file_data)

        zip_buffer.seek(0)
        zip_bytes = zip_buffer.getvalue()

        logger.info(f"[BULK DEBUG] final zip size = {len(zip_bytes)}")

        return zip_bytes, f"{session_id}.zip"

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error in report_bulk_download: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to bulk download reports: {str(e)}")

async def report_download(session_id: str, ens_id: str, file_name: str, type_of_file: str):
    try:
        print("check 1")
        storage_url = get_settings().storage.storage_account_url
        container_name = get_settings().storage.container_name
        sas_token = str(get_settings().storage.sas_token)
        print("check 2")

        blob_service_client = BlobServiceClient(
            account_url=storage_url,
            credential=sas_token
        )
        container_client = blob_service_client.get_container_client(container_name)

        logger.info(f"[DEBUG] container_name = {container_name}")
        logger.info(f"[DEBUG] session_id = {session_id}")
        logger.info(f"[DEBUG] ens_id = {ens_id}")
        logger.info(f"[DEBUG] file_name = {file_name}")
        logger.info(f"[DEBUG] type_of_file = {type_of_file}")

        logger.info("[DEBUG] Listing blobs under session prefix")
        session_prefix = f"{session_id}/"
        session_blobs = list(container_client.list_blobs(name_starts_with=session_prefix))
        for blob in session_blobs:
            logger.info(f"[DEBUG][SESSION] {blob.name}")

        logger.info("[DEBUG] Listing blobs under ens prefix")
        ens_prefix = f"{session_id}/{ens_id}/"
        ens_blobs = list(container_client.list_blobs(name_starts_with=ens_prefix))
        for blob in ens_blobs:
            logger.info(f"[DEBUG][ENS] {blob.name}")

        blob_name = f"{session_id}/{ens_id}/{file_name}.{type_of_file}"
        logger.info(f"[DEBUG] exact blob_name = {blob_name}")

        blob_client = blob_service_client.get_blob_client(
            container=container_name,
            blob=blob_name
        )

        stream = blob_client.download_blob()
        file_data = stream.readall()

        logger.info(f"[DEBUG] Downloaded bytes = {len(file_data)}")
        logger.info(f"[DEBUG] First 8 bytes = {file_data[:8]}")

        return file_data, f"{file_name}.{type_of_file}"

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error in report_download: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to download report: {str(e)}")

# Initialize R2 S3 client
import urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

def get_r2_client():
    from botocore.config import Config
    return boto3.client(
        "s3",
        aws_access_key_id=get_settings().r2_storage.access_key,
        aws_secret_access_key=get_settings().r2_storage.secreate_account_key,
        endpoint_url=get_settings().r2_storage.storage_account_url,
        region_name="auto",
        verify=False,  # temporary until corporate CA is sorted
        config=Config(
            signature_version="s3v4",
            s3={"addressing_style": "path"},  # R2 requires path-style
        )
    )


async def r2_report_bulk_download(session_id: str) -> Tuple:
    try:
        bucket_name = get_settings().r2_storage.storage_container_name
        s3 = get_r2_client()

        prefix = f"{session_id}/"
        response = s3.list_objects_v2(Bucket=bucket_name, Prefix=prefix)

        if "Contents" not in response or not response["Contents"]:
            return None, f"No files found for session_id {session_id}"

        blob_list = response["Contents"]

        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            for obj in blob_list:
                key = obj["Key"]
                if key.endswith("/"):  # Skip folders
                    continue
                file_obj = s3.get_object(Bucket=bucket_name, Key=key)
                file_data = file_obj["Body"].read()

                # Strip "session_id/" prefix from the key
                relative_path = key[len(prefix):]
                zip_file.writestr(relative_path, file_data)

        zip_buffer.seek(0)
        return zip_buffer.getvalue(), f"{session_id}.zip"

    except Exception as e:
        logger.error(f"Error in r2_report_bulk_download: {str(e)}")
        return None, str(e)  # Always return a tuple — never a plain dict


async def local_report_bulk_download(session_id: str) -> Tuple:
    try:
        # Root folder where reports are stored locally
        REPORTS_ROOT = r"C:\ENS\personal repo\ens-engine\src\utils\local-reports"  # ⚠️ change this

        session_path = os.path.join(REPORTS_ROOT, session_id)

        if not os.path.exists(session_path) or not os.path.isdir(session_path):
            return None, f"No files found for session_id {session_id}"

        zip_buffer = io.BytesIO()

        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            for root, dirs, files in os.walk(session_path):
                for file in files:
                    full_path = os.path.join(root, file)

                    # Create relative path inside zip
                    relative_path = os.path.relpath(full_path, session_path)

                    zip_file.write(full_path, arcname=relative_path)

        zip_buffer.seek(0)
        return zip_buffer.getvalue(), f"{session_id}.zip"

    except Exception as e:
        logger.error(f"Error in local_report_bulk_download: {str(e)}")
        return None, str(e)


async def r2_report_download(session_id: str, ens_id: str, type_of_file: str) -> Dict:
    try:
        bucket_name = get_settings().r2_storage.storage_container_name
        s3 = get_r2_client()

        folder_path = f"{session_id}/{ens_id}/"
        response = s3.list_objects_v2(Bucket=bucket_name, Prefix=folder_path)

        if "Contents" not in response:
            raise HTTPException(status_code=404, detail=f"No files found in R2 for {ens_id}")

        matching_files = [
            obj for obj in response["Contents"]
            if obj["Key"].endswith(f".{type_of_file}")
        ]

        if not matching_files:
            raise HTTPException(
                status_code=404,
                detail=f"No matching {type_of_file} file found for {ens_id}"
            )

        latest_file = max(matching_files, key=lambda x: x["LastModified"])
        latest_file_key = latest_file["Key"]

        decoded_filename = urllib.parse.unquote(os.path.basename(latest_file_key))

        file_obj = s3.get_object(Bucket=bucket_name, Key=latest_file_key)
        file_data = file_obj["Body"].read()

        return file_data, decoded_filename

    except ClientError as e:
        logger.error(f"Error in report_download (R2): {str(e)}")
        return {"error": str(e)}


async def local_report_download(session_id: str, ens_id: str, type_of_file: str):
    try:
        # Root directory where reports are stored locally
        REPORTS_ROOT = r"C:\ENS\personal repo\ens-engine\src\utils\local-reports"  # ⚠️ change this

        folder_path = os.path.join(REPORTS_ROOT, session_id, ens_id)

        if not os.path.exists(folder_path) or not os.path.isdir(folder_path):
            raise HTTPException(
                status_code=404,
                detail=f"No files found locally for {ens_id}"
            )

        # Find matching files
        matching_files = [
            os.path.join(folder_path, f)
            for f in os.listdir(folder_path)
            if f.endswith(f".{type_of_file}")
            and os.path.isfile(os.path.join(folder_path, f))
        ]

        if not matching_files:
            raise HTTPException(
                status_code=404,
                detail=f"No matching {type_of_file} file found for {ens_id}"
            )

        # Pick latest file by modified time (same idea as LastModified)
        latest_file_path = max(matching_files, key=os.path.getmtime)

        decoded_filename = urllib.parse.unquote(
            os.path.basename(latest_file_path)
        )

        with open(latest_file_path, "rb") as f:
            file_data = f.read()

        return file_data, decoded_filename

    except Exception as e:
        logger.error(f"Error in report_download (local): {str(e)}")
        return {"error": str(e)}



async def r2_screener_report_bulk_download(session_ids: List[str]) -> Tuple[bytes, str]:
    try:
        bucket_name = get_settings().r2_storage.storage_container_name
        s3 = get_r2_client()

        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            for session_id in session_ids:
                prefix = f"{session_id}/"
                response = s3.list_objects_v2(Bucket=bucket_name, Prefix=prefix)

                if "Contents" not in response or not response["Contents"]:
                    logger.warning(f"No files found for session_id {session_id}")
                    continue

                for obj in response["Contents"]:
                    key = obj["Key"]

                    if key.endswith("/") or key.count("/") < 2:
                        continue

                    relative_key = key[len(prefix):]
                    ens_id = relative_key.split("/", 1)[0]

                    if not ens_id:
                        continue

                    file_obj = s3.get_object(Bucket=bucket_name, Key=key)
                    file_data = file_obj["Body"].read()

                    zip_path = f"{ens_id}/{relative_key.split('/', 1)[1]}"
                    zip_file.writestr(zip_path, file_data)

        zip_buffer.seek(0)
        return zip_buffer.getvalue(), "all-reports.zip"

    except ClientError as e:
        logger.error(f"Error in report_bulk_download (R2): {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to download reports")


DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _parse_date_only_utc(d: str) -> datetime:
    if not DATE_RE.match(d):
        raise ValueError("Date must be in YYYY-MM-DD format")
    return datetime.strptime(d, "%Y-%m-%d").replace(tzinfo=timezone.utc)


async def download_notification_csv_(
    *, session, start_date: str, end_date: str, notificationtypes: list
) -> StreamingResponse:
    try:
        from_ts = _parse_date_only_utc(start_date)
        to_ts_exclusive = _parse_date_only_utc(end_date) + timedelta(days=1)
        if to_ts_exclusive <= from_ts:
            raise ValueError("'end date' must be on/after 'start date'")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    notification = Base.metadata.tables.get("notification")
    entity_universe = Base.metadata.tables.get("entity_universe")
    if notification is None or entity_universe is None:
        raise HTTPException(status_code=500, detail="Missing table(s).")

    filters = [
        notification.c.create_time >= from_ts,
        notification.c.create_time < to_ts_exclusive,
    ]

    if notificationtypes:
        filters.append(notification.c.notification_type.in_(list(dict.fromkeys(notificationtypes))))

    query = (
        select(
            entity_universe.c.external_vendor_id.label("ID"),
            notification.c.ens_id.label("ENS ID"),
            notification.c.notification_type.label("Notification Type"),
            notification.c.title.label("Title"),
            notification.c.description.label("Description"),
            notification.c.theme.label("Theme"),
            entity_universe.c.name.label("Company Name"),
            entity_universe.c.national_id.label("National ID"),
            entity_universe.c.country.label("Country"),
            notification.c.create_time.label("Create Time"),
        )
        .select_from(
            notification.join(
                entity_universe,
                entity_universe.c.ens_id == notification.c.ens_id
            )
        )
        .where(and_(*filters))
        .order_by(notification.c.create_time.desc())
    )

    result = await session.execute(query)
    columns = list(result.keys())
    rows = result.all()
    formatted_res = [dict(zip(columns, row)) for row in rows]

    for d in formatted_res:
        ct = d.get("Create Time")
        if isinstance(ct, datetime):
            if ct.tzinfo:
                ct = ct.astimezone(timezone.utc)
            d["Create Time"] = ct.strftime("%Y-%m-%d %H:%M:%S")

    text_buf = io.StringIO(newline="")
    writer = csv.DictWriter(text_buf, fieldnames=columns)
    writer.writeheader()
    writer.writerows(formatted_res)

    byte_buf = io.BytesIO()
    byte_buf.write("\ufeff".encode("utf-8"))
    byte_buf.write(text_buf.getvalue().encode("utf-8"))
    byte_buf.seek(0)

    return StreamingResponse(
        byte_buf,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="notifications_{start_date}_to_{end_date}.csv"'},
    )


# ---- helper to list ALL objects under a prefix (handles pagination) ----
def _list_all_objects(s3, bucket: str, prefix: str):
    token = None
    while True:
        kwargs = {"Bucket": bucket, "Prefix": prefix}
        if token:
            kwargs["ContinuationToken"] = token
        resp = s3.list_objects_v2(**kwargs)
        for obj in resp.get("Contents", []):
            if obj["Key"].endswith("/"):
                continue
            yield obj
        if not resp.get("IsTruncated"):
            break
        token = resp.get("NextContinuationToken")


# ---- DB step: get (session_id, ens_id) pairs for given group_id ----
async def get_session_ens_pairs(
    *, session: AsyncSession, group_id: str
) -> List[Tuple[str, str]]:
    md = Base.metadata
    esgm = md.tables.get("ens_schedule_group_mapping")
    eu = md.tables.get("entity_universe")

    if esgm is None or eu is None:
        raise HTTPException(
            status_code=404,
            detail="Table 'ens_schedule_group_mapping' or 'entity_universe' does not exist in the database schema."
        )

    pairs_stmt = (
        select(
            eu.c.last_session_id.label("session_id"),
            eu.c.ens_id.label("ens_id"),
        )
        .select_from(
            esgm.join(eu, eu.c.ens_id == esgm.c.ens_id)
        )
        .where(
            esgm.c.group_id == group_id,
            eu.c.last_session_id.isnot(None),
        )
        .distinct()
    )

    res = await session.execute(pairs_stmt)
    pairs = {(r.session_id, r.ens_id) for r in res}
    return pairs


# ---- R2 bulk download for a given group_id (no file-type filter) ----
async def r2_screener_report_bulk_download_by_source(
    session: AsyncSession, group_id: str
) -> Tuple[bytes, str]:
    try:
        pairs = await get_session_ens_pairs(session=session, group_id=group_id)
        if not pairs:
            raise HTTPException(status_code=404, detail="No (session_id, ens_id) pairs found for this group_id")

        bucket_name = get_settings().r2_storage.storage_container_name
        s3 = get_r2_client()

        any_files = False
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zipf:
            for session_id, ens_id in pairs:
                prefix = f"{session_id}/{ens_id}/"
                try:
                    for obj in _list_all_objects(s3, bucket_name, prefix):
                        any_files = True
                        key = obj["Key"]
                        relative_under_ens = key[len(prefix):]
                        zip_path = f"{ens_id}/{relative_under_ens}"
                        file_obj = s3.get_object(Bucket=bucket_name, Key=key)
                        zipf.writestr(zip_path, file_obj["Body"].read())
                except ClientError as ce:
                    logger.warning(f"R2 list/get failed for prefix {prefix}: {ce}")
                    continue

        if not any_files:
            raise HTTPException(status_code=404, detail="No files found in R2 for the resolved session/entity pairs")

        zip_buffer.seek(0)
        return zip_buffer.getvalue(), f"reports_{group_id}.zip"

    except ClientError as e:
        logger.error(f"Error in report_bulk_download_by_source (R2): {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to download reports")

import base64  # add this at the top with other imports


async def get_google_images_from_blob(google_image_name: str) -> List[Dict]:
    """
    Fetch all images from the 'images' Azure Blob container whose blob names
    start with the given google_image_name prefix.

    Returns a list of dicts each with:
      - filename:     base filename of the blob
      - content_type: inferred MIME type (e.g. "image/png")
      - data:         base64-encoded image bytes

    Frontend usage:
      <img src={`data:${item.content_type};base64,${item.data}`} />
    """
    try:
        storage_url = get_settings().storage.storage_account_url
        images_container_name = get_settings().storage.images_container_name
        images_sas_token = str(get_settings().storage.images_sas_token)

        blob_service_client = BlobServiceClient(
            account_url=storage_url,
            credential=images_sas_token
        )
        container_client = blob_service_client.get_container_client(images_container_name)

        logger.info(f"[IMAGE DEBUG] container_name = {images_container_name}")
        logger.info(f"[IMAGE DEBUG] google_image_name prefix = {google_image_name}")

        blob_list = list(container_client.list_blobs(name_starts_with=google_image_name))

        logger.info(f"[IMAGE DEBUG] blob count = {len(blob_list)}")

        if not blob_list:
            raise HTTPException(
                status_code=404,
                detail=f"No images found with prefix '{google_image_name}'"
            )

        content_type_map = {
            "jpg":  "image/jpeg",
            "jpeg": "image/jpeg",
            "png":  "image/png",
            "gif":  "image/gif",
            "webp": "image/webp",
            "svg":  "image/svg+xml",
            "bmp":  "image/bmp",
        }

        images = []
        for blob in blob_list:
            blob_name = blob.name

            if blob_name.endswith("/"):
                logger.info(f"[IMAGE DEBUG] skipping folder blob = {blob_name}")
                continue

            logger.info(f"[IMAGE DEBUG] downloading blob = {blob_name}")

            blob_client = blob_service_client.get_blob_client(
                container=images_container_name,
                blob=blob_name
            )
            file_data = blob_client.download_blob().readall()

            ext = os.path.splitext(blob_name)[1].lower().lstrip(".")
            content_type = content_type_map.get(ext, "image/jpeg")

            encoded = base64.b64encode(file_data).decode("utf-8")

            images.append({
                "filename":     os.path.basename(blob_name),
                "content_type": content_type,
                "data":         encoded,
            })

            logger.info(f"[IMAGE DEBUG] processed image = {blob_name}, size = {len(file_data)} bytes")

        return images

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error in get_google_images_from_blob: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to fetch images: {str(e)}")